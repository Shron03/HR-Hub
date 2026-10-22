import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_hr_hub_presentation():
    prs = Presentation()
    # 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette - Clean, Light, Professional
    BG_COLOR = RGBColor(248, 250, 252)       # Soft off-white / slate-50
    CARD_BG = RGBColor(255, 255, 255)        # Pure white
    CARD_BORDER = RGBColor(226, 232, 240)    # Light slate-200
    TEXT_MAIN = RGBColor(15, 23, 42)         # Deep slate-900
    TEXT_MUTED = RGBColor(100, 116, 139)     # Slate-500
    TEXT_BODY = RGBColor(51, 65, 85)         # Slate-700
    PRIMARY_BLUE = RGBColor(37, 99, 235)     # Blue-600
    ACCENT_LIGHT = RGBColor(239, 246, 255)   # Blue-50
    ACCENT_BORDER = RGBColor(191, 219, 254)  # Blue-200
    ACCENT_TEXT = RGBColor(29, 78, 216)      # Blue-700

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()

    def add_header(slide, tag_text, title_text, subtitle_text, page_number):
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(3.2), Inches(0.32))
        pill.fill.solid()
        pill.fill.fore_color.rgb = ACCENT_LIGHT
        pill.line.color.rgb = ACCENT_BORDER
        pill.line.width = Pt(1)
        tf_pill = pill.text_frame
        tf_pill.word_wrap = True
        tf_pill.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_pill = tf_pill.paragraphs[0]
        p_pill.text = tag_text.upper()
        p_pill.font.size = Pt(9.5)
        p_pill.font.bold = True
        p_pill.font.color.rgb = ACCENT_TEXT
        p_pill.alignment = PP_ALIGN.CENTER

        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(10.5), Inches(0.6))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_MAIN

        tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(10.5), Inches(0.4))
        tf_sub = tb_sub.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle_text
        p_sub.font.size = Pt(11.5)
        p_sub.font.color.rgb = TEXT_MUTED

        tb_num = slide.shapes.add_textbox(Inches(11.5), Inches(0.45), Inches(1.0), Inches(0.4))
        tf_num = tb_num.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.text = f"Slide {page_number} of 4"
        p_num.font.size = Pt(10)
        p_num.font.bold = True
        p_num.font.color.rgb = TEXT_MUTED
        p_num.alignment = PP_ALIGN.RIGHT

    def create_card(slide, left, top, width, height, title="", subtitle=""):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.05)
        tf.margin_right = Inches(0.05)
        tf.margin_top = Inches(0.05)
        tf.margin_bottom = Inches(0.05)

        if title:
            p_t = tf.paragraphs[0]
            p_t.text = title
            p_t.font.size = Pt(13)
            p_t.font.bold = True
            p_t.font.color.rgb = TEXT_MAIN
            p_t.space_after = Pt(2)

            if subtitle:
                p_s = tf.add_paragraph()
                p_s.text = subtitle
                p_s.font.size = Pt(9.5)
                p_s.font.color.rgb = TEXT_MUTED
                p_s.space_after = Pt(8)
            else:
                p_t.space_after = Pt(8)

        return tf

    # ==========================================
    # SLIDE 1: INTRODUCTION & PROJECT OVERVIEW
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)
    add_header(
        slide1,
        "Project Introduction & Architecture",
        "HR Hub Enterprise - Workforce Operations Platform",
        "Full-stack human resource management, live talent analytics, and dynamic KPI calibration system",
        1
    )

    card_w3 = Inches(3.64)
    card_h = Inches(5.1)
    top_pos = Inches(1.85)

    # Slide 1 - Card 1: Project Overview
    tf1 = create_card(slide1, Inches(0.8), top_pos, card_w3, card_h, "Project Overview & Vision", "Modernizing corporate HR operations")
    items1 = [
        ("Direct Landing Experience:", " Direct entry into the executive HR Director panel on initial site visit without mandatory login roadblocks."),
        ("Enterprise Architecture:", " Built to match Zoho People and Google Workspace standards with role-specific views and modular workflows."),
        ("Widescreen Sticky Layout:", " Replaced sidebars with a sticky top navigation menu, maximizing horizontal desktop space."),
        ("Real-Time KPI Calibration:", " Dynamic deliverable tracking, decimal performance ratings (0.0-10.0), and live composite score derivations."),
        ("Indian Workforce Dataset:", " Pre-configured with a 15-member Indian corporate roster, reporting lines, compensation CTC, and departments.")
    ]
    for bold_prefix, text in items1:
        p = tf1.add_paragraph()
        p.space_after = Pt(6)
        r1 = p.add_run()
        r1.text = "• " + bold_prefix
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = PRIMARY_BLUE
        r2 = p.add_run()
        r2.text = text
        r2.font.bold = False
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_BODY

    # Slide 1 - Card 2: Problem Statement & Need
    tf2 = create_card(slide1, Inches(4.84), top_pos, card_w3, card_h, "Problem Statement & Need", "Challenges in traditional HR platforms")
    items2 = [
        ("Fragmented Tooling:", " Legacy HR operations rely on disconnected spreadsheets and manual email approvals, causing operational latency."),
        ("Rigid Role Boundaries:", " Off-the-shelf software often lacks granular Role-Based Access Control (RBAC), exposing sensitive hierarchy data."),
        ("Opaque Evaluation Metrics:", " Performance grading is often calculated via hidden static annual formulas rather than transparent deliverables."),
        ("Cluttered User Interfaces:", " Enterprise portals suffer from tall, scroll-heavy sidebars and wasted desktop margins, diminishing daily usability."),
        ("Absence of Instant Recovery:", " Lack of standardized demonstration datasets and instant persistence resets complicates stakeholder evaluation.")
    ]
    for bold_prefix, text in items2:
        p = tf2.add_paragraph()
        p.space_after = Pt(6)
        r1 = p.add_run()
        r1.text = "• " + bold_prefix
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = PRIMARY_BLUE
        r2 = p.add_run()
        r2.text = text
        r2.font.bold = False
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_BODY

    # Slide 1 - Card 3: Core Architecture & Stack
    tf3 = create_card(slide1, Inches(8.88), top_pos, card_w3, card_h, "System Architecture & Stack", "Clean decoupled client-server model")
    items3 = [
        ("Frontend Application:", " Built with React 18, Vite 8 build tool, and Tailwind CSS for a modern, responsive user experience."),
        ("Data Visualization:", " Integrated Chart.js and React-Chartjs-2 for radar charts, line trends, and distribution doughnuts."),
        ("Backend Server & API:", " Node.js runtime and Express.js REST API with modular controllers, routes, and middleware."),
        ("Database Engine:", " MongoDB with Mongoose ORM, featuring schema validation, relationships, and TTL auto-expiry indexes."),
        ("Enterprise RBAC:", " Decoupled role routing for HR Director, Engineering Manager, and Employee Self-Service."),
        ("Data Persistence:", " All custom records, leave states, and calibrated CTCs persist in the database across page refreshes.")
    ]
    for bold_prefix, text in items3:
        p = tf3.add_paragraph()
        p.space_after = Pt(5)
        r1 = p.add_run()
        r1.text = "• " + bold_prefix
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = PRIMARY_BLUE
        r2 = p.add_run()
        r2.text = text
        r2.font.bold = False
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_BODY


    # ==========================================
    # SLIDE 2: WEBPAGE WORKING & CORE FEATURES
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(
        slide2,
        "Webpage Working & Features",
        "Operational Workflows & Core System Features",
        "Modular architecture partitioned across HR Director, Engineering Manager, and Employee workspaces",
        2
    )

    grid_w = Inches(5.68)
    grid_h = Inches(2.45)
    row1_top = Inches(1.85)
    row2_top = Inches(4.5)

    # Slide 2 - Box 1: HR Executive Modules
    tf_w1 = create_card(slide2, Inches(0.8), row1_top, grid_w, grid_h, "HR Director: Governance & Operations", "7 comprehensive management modules")
    w1_items = [
        ("Organization & Employee Directory:", " Unified personnel directory and 3-in-1 organizational reporting tree with granular profile editing."),
        ("Onboarding & Offboarding (HR Process):", " Interactive checklist tracking for documentation, IT assets provisioning, and exit clearances."),
        ("Time & Attendance Administration:", " Presence compliance monitor and leave authorization queue with start/end date constraints."),
        ("Talent Management & Appraisals:", " Real-time KPI simulator, decimal reviews (0.0-10.0), 9-box talent matrix, and appraisal dispatch.")
    ]
    for bold_prefix, text in w1_items:
        p = tf_w1.add_paragraph()
        p.space_after = Pt(2.5)
        r1 = p.add_run()
        r1.text = "• " + bold_prefix
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = PRIMARY_BLUE
        r2 = p.add_run()
        r2.text = text
        r2.font.size = Pt(9)
        r2.font.color.rgb = TEXT_BODY

    # Slide 2 - Box 2: HR Services & Documents
    tf_w2 = create_card(slide2, Inches(6.84), row1_top, grid_w, grid_h, "HR Director: Services & Intelligence", "Employee relations and executive telemetry")
    w2_items = [
        ("HR Help Desk & Requests:", " Triage employee support tickets (Payroll, Leave, IT, Policies) with tracked resolution notes."),
        ("Files & Document Management:", " Enterprise repository for company policies, POSH handbooks, NDA templates, and tax declarations."),
        ("Analytics & Reports:", " Executive operations cockpit, 6-month attendance trajectory, strategic insights, and CSV exports."),
        ("Reset Demo Capability:", " Persistent database modifications with a 1-click Reset Demo button to restore standard data.")
    ]
    for bold_prefix, text in w2_items:
        p = tf_w2.add_paragraph()
        p.space_after = Pt(2.5)
        r1 = p.add_run()
        r1.text = "• " + bold_prefix
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = PRIMARY_BLUE
        r2 = p.add_run()
        r2.text = text
        r2.font.size = Pt(9)
        r2.font.color.rgb = TEXT_BODY

    # Slide 2 - Box 3: Engineering Manager Modules
    tf_w3 = create_card(slide2, Inches(0.8), row2_top, grid_w, grid_h, "Engineering Manager: Sprint Oversight", "4 dedicated managerial supervisory modules")
    w3_items = [
        ("Approval Center:", " Timesheet hour sign-offs and team leave review (monitoring presence while HR retains final approval authority)."),
        ("Team Supervision:", " Dedicated engineering roster oversight, live presence monitoring (On-site, Hybrid, Remote), and task delegation."),
        ("Performance & Goal Tracking:", " Team sprint velocity, overall team KPI benchmark, decimal review calibrator, and quarterly OKRs."),
        ("Time Logs & Projects:", " Project deliverable boards, team hours burn-down, sprint completion rates, and blocker resolution.")
    ]
    for bold_prefix, text in w3_items:
        p = tf_w3.add_paragraph()
        p.space_after = Pt(2.5)
        r1 = p.add_run()
        r1.text = "• " + bold_prefix
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = PRIMARY_BLUE
        r2 = p.add_run()
        r2.text = text
        r2.font.size = Pt(9)
        r2.font.color.rgb = TEXT_BODY

    # Slide 2 - Box 4: Employee Self-Service Modules
    tf_w4 = create_card(slide2, Inches(6.84), row2_top, grid_w, grid_h, "Employee Self-Service Portal", "5 specialized employee self-service modules")
    w4_items = [
        ("Self Attendance & Leave:", " Live punch clock-in/out simulator, leave balances (CL, SL, PL), and time-off request submission."),
        ("Time Tracking:", " Daily timesheet logger, logged hours history, and active sprint deliverables completion toggle."),
        ("Performance & Self-Appraisal:", " Personal KPI scorecard, quarterly self-evaluation submitter, and official HR appraisals timeline."),
        ("Profile & Payroll Documents:", " Digital employee badge, annual CTC breakdown (INR), and simulated monthly salary slip generator."),
        ("HR Tickets:", " Raise confidential support inquiries to HR with live status badges (Open, In Progress, Resolved).")
    ]
    for bold_prefix, text in w4_items:
        p = tf_w4.add_paragraph()
        p.space_after = Pt(2.5)
        r1 = p.add_run()
        r1.text = "• " + bold_prefix
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = PRIMARY_BLUE
        r2 = p.add_run()
        r2.text = text
        r2.font.size = Pt(9)
        r2.font.color.rgb = TEXT_BODY


    # ==========================================
    # SLIDE 3: IDEAS IMPLEMENTED & TOOLS USED
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(
        slide3,
        "Implementation & Tooling",
        "Ideas Implemented, Engineering Innovations & Tools Used",
        "Technical concepts realized in code along with the complete engineering toolkit",
        3
    )

    card_w2 = Inches(5.68)
    card_h2 = Inches(5.1)

    # Slide 3 - Left Column: Key Ideas Implemented
    tf_i1 = create_card(slide3, Inches(0.8), top_pos, card_w2, card_h2, "Key Ideas & Innovations Implemented", "Strategic concepts translated into platform features")
    ideas_items = [
        ("Direct First-Time Landing:", " Automatic initialization into the HR Director dashboard on initial load, bypassing login barriers."),
        ("Two-Phase Administrative Verification:", " Formal HR registration queue preventing unapproved users from logging in until vetted."),
        ("Widescreen Sticky Navigation Paradigm:", " Transformed the traditional cramped sidebar into a horizontal sticky top bar, reclaiming 15% screen space."),
        ("Real-Time Deliverable Synchronization:", " Implemented bidirectional constraints between tasks assigned and tasks completed."),
        ("Contextual Role Privacy Partitioning:", " Prevented unauthorized hierarchy browsing by replacing the org chart with personal self-service views."),
        ("Time-Bound TTL Notification Lifecycle:", " Combined MongoDB native TTL indexes (259,200 seconds) with API date filters for a 3-day window."),
        ("Zero-Formula Presentation Layer:", " Abstracted complex arithmetic calculations into executive visual indicators."),
        ("End-to-End Database Persistence:", " Tickets, onboarding checklists, documents, and time logs persist across reloads.")
    ]
    for bold_prefix, text in ideas_items:
        p = tf_i1.add_paragraph()
        p.space_after = Pt(3.5)
        r1 = p.add_run()
        r1.text = "• " + bold_prefix
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = PRIMARY_BLUE
        r2 = p.add_run()
        r2.text = text
        r2.font.size = Pt(9)
        r2.font.color.rgb = TEXT_BODY

    # Slide 3 - Right Column: Tools & Technologies Used
    tf_i2 = create_card(slide3, Inches(6.84), top_pos, card_w2, card_h2, "Tools, Libraries & Technologies Used", "End-to-end modern web technology stack")
    tools_items = [
        ("React 18 & Vite 8:", " Client-side component hierarchy, high-performance Virtual DOM updates, and fast Hot Module Replacement."),
        ("Tailwind CSS:", " Modern utility-first CSS framework enabling clean layout management and frosted glassmorphic card effects."),
        ("Chart.js & React-Chartjs-2:", " Rendered HTML5 canvas charts including radar skill benchmarks, 6-month attendance lines, and doughnuts."),
        ("Lucide React Icon Library:", " Over 30 standardized enterprise icons providing visual clarity without using informal emojis."),
        ("Node.js & Express.js:", " Asynchronous runtime and RESTful API framework handling authentication, records, and workflows."),
        ("MongoDB & Mongoose ODM:", " Document database utilizing structured schemas, relational lookups, and TTL index auto-deletion."),
        ("JSON Web Tokens (JWT):", " Secure stateless authorization tokens verifying role claims across protected routes."),
        ("Git & GitHub:", " Version control repository maintaining point-by-point commit history ([github.com/Shron03/HR-Hub](https://github.com/Shron03/HR-Hub)).")
    ]
    for bold_prefix, text in tools_items:
        p = tf_i2.add_paragraph()
        p.space_after = Pt(3.5)
        r1 = p.add_run()
        r1.text = "• " + bold_prefix
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = PRIMARY_BLUE
        r2 = p.add_run()
        r2.text = text
        r2.font.size = Pt(9)
        r2.font.color.rgb = TEXT_BODY


    # ==========================================
    # SLIDE 4: LEARNING OUTCOMES & REFERENCES
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(
        slide4,
        "Learning Outcomes & References",
        "Project Learning Outcomes & Benchmark References",
        "Core competencies gained throughout development alongside industry references and citations",
        4
    )

    # Slide 4 - Left Column: Learning Outcomes
    tf_l1 = create_card(slide4, Inches(0.8), top_pos, card_w2, card_h2, "Key Project Learning Outcomes", "Skills and technical competencies acquired")
    outcomes_items = [
        ("Full-Stack Enterprise Architecture:", " Mastered end-to-end coordination between client-side UI states, Express route controllers, and MongoDB schemas."),
        ("Role-Based Security & Authorization (RBAC):", " Implemented multi-tier user privilege separation across both frontend views and backend API endpoints."),
        ("Responsive Layout & Viewport Optimization:", " Acquired deep knowledge in widescreen layout design, replacing vertical sidebars with sticky top headers."),
        ("Real-Time KPI Calibration Systems:", " Developed interactive calculation engines synchronizing interdependent variables without performance lag."),
        ("Automated Multi-Stage Test Engineering:", " Designed and executed 6 automated test suites covering boundary math, concurrency, and authorization gates."),
        ("Enterprise UI/UX Design Standards:", " Crafted an uncluttered, accessible user interface following professional corporate design principles.")
    ]
    for bold_prefix, text in outcomes_items:
        p = tf_l1.add_paragraph()
        p.space_after = Pt(5)
        r1 = p.add_run()
        r1.text = "• " + bold_prefix
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = PRIMARY_BLUE
        r2 = p.add_run()
        r2.text = text
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_BODY

    # Slide 4 - Right Column: References & Citations
    tf_l2 = create_card(slide4, Inches(6.84), top_pos, card_w2, card_h2, "References & Benchmark Sources", "Design inspirations, technical documentation & source code")
    refs_items = [
        ("Zoho People Enterprise Portal:", " Benchmark reference for leave request authorization workflows, attendance compliance, and employee self-service modules."),
        ("Google Workspace & Cloud Console:", " Visual reference for clean typography, sticky horizontal navigation bars, and uncluttered desktop data cards."),
        ("Official React.js & Vite Documentation:", " Component lifecycle architecture, Virtual DOM optimization patterns, and production bundling guidelines."),
        ("Tailwind CSS Design Framework:", " Utility-first responsive design principles, color theory, and subtle glassmorphic styling techniques."),
        ("Human Resources Capstone Project Specifications:", " Capstone curriculum guidelines covering workforce metrics, leave tracking, and organizational tree hierarchies."),
        ("Project Source Code Repository:", " Public GitHub Repository maintaining full commit history, documentation, and source code:"),
        ("Repository Link:", " https://github.com/Shron03/HR-Hub")
    ]
    for bold_prefix, text in refs_items:
        p = tf_l2.add_paragraph()
        p.space_after = Pt(5)
        r1 = p.add_run()
        r1.text = "• " + bold_prefix
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = PRIMARY_BLUE
        r2 = p.add_run()
        r2.text = text
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_BODY

    # Save Presentation to Workspace
    output_path = r"C:\Users\Aethon\OneDrive\Desktop\Internship\HR_Hub_Project_Presentation.pptx"
    prs.save(output_path)
    print(f"[SUCCESS] Presentation generated successfully at: {output_path}")

    # Also save a copy to root user directory
    root_copy = r"C:\Users\Aethon\HR_Hub_Project_Presentation.pptx"
    prs.save(root_copy)
    print(f"[SUCCESS] Secondary copy saved at: {root_copy}")

if __name__ == "__main__":
    create_hr_hub_presentation()
